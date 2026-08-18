from __future__ import annotations

import csv
from io import StringIO
from pathlib import Path

from openpyxl import load_workbook

from preprocess_file.handlers.base import Handler
from preprocess_file.handlers.text import decode_bytes
from preprocess_file.models import Document, Element, ElementType, FileKind, ProcessContext
from preprocess_file.output import table_to_markdown


class SpreadsheetHandler(Handler):
    kinds = (FileKind.XLSX, FileKind.XLS, FileKind.CSV, FileKind.TSV)

    def process(self, ctx: ProcessContext) -> Document:
        kind = ctx.kind or FileKind.XLSX
        document = Document(source=str(ctx.path), file_type=kind)
        if kind in {FileKind.CSV, FileKind.TSV}:
            _load_delimited(ctx.path, document, delimiter="," if kind is FileKind.CSV else "\t")
            return document
        if kind is FileKind.XLS:
            document.warnings.append("Legacy .xls is not parsed natively; convert to .xlsx first.")
            return document
        _load_xlsx(ctx, document)
        return document


def _load_delimited(path: Path, document: Document, delimiter: str) -> None:
    raw = path.read_bytes()
    text, encoding = decode_bytes(raw)
    sample = text[:4096]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=delimiter + ",\t;|")
        delimiter = dialect.delimiter
    except csv.Error:
        pass
    rows = list(csv.reader(StringIO(text), delimiter=delimiter))
    header_row = detect_header_row(rows)
    table_rows = rows[header_row:]
    markdown = table_to_markdown(_stringify(table_rows))
    document.metadata.update(
        {
            "encoding": encoding,
            "delimiter": delimiter,
            "header_row": header_row,
            "rows": max(len(table_rows) - 1, 0),
        }
    )
    if markdown:
        document.add(
            Element(
                type=ElementType.TABLE,
                text=markdown,
                metadata={"header_row": header_row, "sheet": path.stem},
            )
        )


def _load_xlsx(ctx: ProcessContext, document: Document) -> None:
    formulas = load_workbook(ctx.path, data_only=False, read_only=False)
    values = load_workbook(ctx.path, data_only=True, read_only=False)
    document.metadata["sheets"] = formulas.sheetnames
    missing_cache = False

    for sheet_name in formulas.sheetnames:
        ws_f = formulas[sheet_name]
        ws_v = values[sheet_name] if sheet_name in values.sheetnames else ws_f
        if ctx.options.spreadsheet_forward_fill:
            _forward_fill_merged(ws_v)
            _forward_fill_merged(ws_f)

        grid: list[list[str]] = []
        for row_idx, row in enumerate(ws_v.iter_rows(values_only=False), start=1):
            rendered: list[str] = []
            for col_idx, cell in enumerate(row, start=1):
                value = cell.value
                if value is None:
                    formula_cell = ws_f.cell(row_idx, col_idx)
                    if isinstance(formula_cell.value, str) and formula_cell.value.startswith("="):
                        missing_cache = True
                        value = formula_cell.value
                rendered.append("" if value is None else str(value))
            grid.append(rendered)

        header_row = detect_header_row(grid)
        table_rows = _drop_leading_empty(grid[header_row:])
        markdown = table_to_markdown(table_rows)
        if not markdown:
            continue
        document.add(
            Element(
                type=ElementType.TABLE,
                text=markdown,
                metadata={
                    "sheet": sheet_name,
                    "header_row": header_row,
                    "rows": max(len(table_rows) - 1, 0),
                },
            )
        )

    if missing_cache:
        document.warnings.append(
            "Some formula cells had no cached value; formula text was kept instead of computed results."
        )


def detect_header_row(rows: list[list[object]], max_scan: int = 20) -> int:
    best_index = 0
    best_score = -1.0
    scan = rows[:max_scan]
    for index, row in enumerate(scan):
        nonempty = [cell for cell in row if str(cell).strip()]
        strings = [cell for cell in nonempty if not _is_number(cell)]
        if len(nonempty) < 2:
            continue
        score = len(strings) * 2 + len(nonempty) - index * 0.1
        if len(strings) >= 2 and score > best_score:
            best_score = score
            best_index = index
    return best_index


def _forward_fill_merged(worksheet) -> None:
    ranges = list(worksheet.merged_cells.ranges)
    snapshots = []
    for rng in ranges:
        snapshots.append((rng, worksheet.cell(rng.min_row, rng.min_col).value))
    for rng, value in snapshots:
        worksheet.unmerge_cells(str(rng))
        for row in range(rng.min_row, rng.max_row + 1):
            for col in range(rng.min_col, rng.max_col + 1):
                if worksheet.cell(row, col).value is None:
                    worksheet.cell(row, col).value = value


def _drop_leading_empty(rows: list[list[str]]) -> list[list[str]]:
    while rows and not any(cell.strip() for cell in rows[0]):
        rows = rows[1:]
    while rows and not any(cell.strip() for cell in rows[-1]):
        rows = rows[:-1]
    return rows


def _stringify(rows: list[list[object]]) -> list[list[str]]:
    return [["" if cell is None else str(cell) for cell in row] for row in rows]


def _is_number(value: object) -> bool:
    if isinstance(value, (int, float)):
        return True
    text = str(value).strip().replace(",", "")
    if not text:
        return False
    try:
        float(text)
        return True
    except ValueError:
        return False

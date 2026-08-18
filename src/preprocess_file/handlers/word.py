from __future__ import annotations

from docx import Document as DocxDocument
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph

from preprocess_file.handlers.base import Handler
from preprocess_file.models import Document, Element, ElementType, FileKind, ProcessContext
from preprocess_file.ocr import ocr_image
from preprocess_file.output import table_to_markdown


class WordHandler(Handler):
    kinds = (FileKind.DOCX, FileKind.DOC, FileKind.RTF)

    def process(self, ctx: ProcessContext) -> Document:
        kind = ctx.kind or FileKind.DOCX
        document = Document(source=str(ctx.path), file_type=kind)
        if kind is not FileKind.DOCX:
            document.warnings.append(
                f"{kind.value} is not parsed natively in v1; convert to .docx (LibreOffice) first."
            )
            return document

        docx = DocxDocument(str(ctx.path))
        document.metadata["paragraphs"] = len(docx.paragraphs)
        document.metadata["tables"] = len(docx.tables)

        for block in _iter_block_items(docx):
            if isinstance(block, Paragraph):
                _add_paragraph(document, block)
            elif isinstance(block, Table):
                rows = [[_cell_text(cell) for cell in row.cells] for row in block.rows]
                markdown = table_to_markdown(rows)
                if markdown:
                    document.add(Element(type=ElementType.TABLE, text=markdown))

        if ctx.options.extract_images:
            _extract_images(docx, document, ctx)
        return document


def _iter_block_items(docx: DocxDocument):
    body = docx.element.body
    for child in body.iterchildren():
        if child.tag == qn("w:p"):
            yield Paragraph(child, docx)
        elif child.tag == qn("w:tbl"):
            yield Table(child, docx)


def _add_paragraph(document: Document, paragraph: Paragraph) -> None:
    text = paragraph.text.strip()
    if not text:
        return
    style = paragraph.style.name if paragraph.style is not None else ""
    if style.startswith("Heading"):
        level = "".join(ch for ch in style if ch.isdigit()) or "1"
        document.add(
            Element(type=ElementType.TITLE, text=text, metadata={"level": int(level), "style": style})
        )
        return
    if style.startswith("List"):
        document.add(Element(type=ElementType.LIST, text=text, metadata={"style": style}))
        return
    document.add(Element(type=ElementType.NARRATIVE, text=text, metadata={"style": style}))


def _cell_text(cell) -> str:
    return " ".join(paragraph.text.strip() for paragraph in cell.paragraphs if paragraph.text.strip())


def _extract_images(docx: DocxDocument, document: Document, ctx: ProcessContext) -> None:
    count = 0
    for rel in docx.part.rels.values():
        if "image" not in rel.reltype:
            continue
        count += 1
        try:
            from io import BytesIO
            from PIL import Image

            blob = rel.target_part.blob
            image = Image.open(BytesIO(blob)).convert("RGB")
        except Exception:
            document.warnings.append(f"Failed to decode embedded image {rel.rId}")
            continue
        text, confidence = ocr_image(image, ctx.options) if ctx.options.ocr_enabled else ("", None)
        document.add(
            Element(
                type=ElementType.IMAGE,
                text=text,
                metadata={"rel": rel.rId, "ocr_confidence": confidence},
            )
        )
    document.metadata["images"] = count

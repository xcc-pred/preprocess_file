from __future__ import annotations

from io import BytesIO

from PIL import Image
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from preprocess_file.handlers.base import Handler
from preprocess_file.models import Document, Element, ElementType, FileKind, ProcessContext
from preprocess_file.ocr import ocr_image
from preprocess_file.output import table_to_markdown


class PptHandler(Handler):
    kinds = (FileKind.PPTX, FileKind.PPT)

    def process(self, ctx: ProcessContext) -> Document:
        kind = ctx.kind or FileKind.PPTX
        document = Document(source=str(ctx.path), file_type=kind)
        if kind is FileKind.PPT:
            document.warnings.append("Legacy .ppt is not parsed natively; convert to .pptx first.")
            return document

        presentation = Presentation(str(ctx.path))
        document.metadata["slides"] = len(presentation.slides)
        for index, slide in enumerate(presentation.slides, start=1):
            if index > 1:
                document.add(Element(type=ElementType.PAGE_BREAK, page=index))
            shapes = sorted(
                slide.shapes,
                key=lambda shape: (int(getattr(shape, "top", 0) or 0), int(getattr(shape, "left", 0) or 0)),
            )
            for shape in shapes:
                _add_shape(document, shape, index, ctx)
            notes = _slide_notes(slide)
            if notes:
                document.add(
                    Element(
                        type=ElementType.NARRATIVE,
                        text=notes,
                        page=index,
                        metadata={"role": "notes"},
                    )
                )
        return document


def _add_shape(document: Document, shape, page: int, ctx: ProcessContext) -> None:
    if shape.has_table:
        rows = [[cell.text.strip() for cell in row.cells] for row in shape.table.rows]
        markdown = table_to_markdown(rows)
        if markdown:
            document.add(Element(type=ElementType.TABLE, text=markdown, page=page))
        return
    if shape.has_text_frame:
        text = "\n".join(
            paragraph.text.strip() for paragraph in shape.text_frame.paragraphs if paragraph.text.strip()
        )
        if text:
            is_title = _is_title_shape(shape)
            document.add(
                Element(
                    type=ElementType.TITLE if is_title else ElementType.NARRATIVE,
                    text=text,
                    page=page,
                    metadata={"level": 2} if is_title else {},
                )
            )
        return
    if ctx.options.extract_images and shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
        try:
            image = Image.open(BytesIO(shape.image.blob)).convert("RGB")
        except Exception:
            document.warnings.append(f"Failed to decode image on slide {page}")
            return
        text, confidence = ocr_image(image, ctx.options) if ctx.options.ocr_enabled else ("", None)
        document.add(
            Element(
                type=ElementType.IMAGE,
                text=text,
                page=page,
                metadata={"ocr_confidence": confidence},
            )
        )


def _is_title_shape(shape) -> bool:
    try:
        if not shape.is_placeholder:
            return False
        return "TITLE" in str(shape.placeholder_format.type).upper()
    except Exception:
        return False


def _slide_notes(slide) -> str:
    if not slide.has_notes_slide:
        return ""
    notes_frame = slide.notes_slide.notes_text_frame
    if notes_frame is None:
        return ""
    return "\n".join(paragraph.text.strip() for paragraph in notes_frame.paragraphs if paragraph.text.strip())
